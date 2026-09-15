"""数据工厂单元测试：不依赖 GPU / Ollama，用合成 Word 文档验证全链路。"""

from __future__ import annotations

from pathlib import Path

import pytest

from apps.data_factory import assemble as asm
from apps.data_factory import chunker, docx_extract, splitter


def _make_docx(tmp_path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("第一章 总则", level=1)
    doc.add_paragraph("本制度适用于公司全体正式员工。")
    doc.add_heading("第二章 考勤管理", level=1)
    doc.add_paragraph("第一条 员工应当遵守工作时间，迟到三十分钟以上按旷工半日处理。")
    doc.add_paragraph("第二条 因故不能到岗的，应当提前一个工作日向部门负责人请假。")
    long_para = "第三条 加班应当事先申请，" + "加班补偿按国家规定执行。" * 30
    doc.add_paragraph(long_para)
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "迟到时长"
    table.cell(0, 1).text = "处理方式"
    table.cell(1, 0).text = "30分钟以内"
    table.cell(1, 1).text = "警告"
    table.cell(2, 0).text = "30分钟以上"
    table.cell(2, 1).text = "旷工半日"
    path = tmp_path / "合成制度.docx"
    doc.save(str(path))
    return path


def test_extract_headings_paras_table(tmp_path: Path) -> None:
    blocks = docx_extract.extract_docx(_make_docx(tmp_path))
    kinds = [b.kind for b in blocks]
    assert "heading" in kinds and "para" in kinds and "table" in kinds
    headings = [b for b in blocks if b.kind == "heading"]
    assert headings[0].level == 1 and "总则" in headings[0].text
    table = next(b for b in blocks if b.kind == "table")
    assert "迟到时长：30分钟以内" in table.text


def test_chunker_sections_and_bounds(tmp_path: Path) -> None:
    blocks = docx_extract.extract_docx(_make_docx(tmp_path))
    chunks = chunker.chunk_blocks(blocks, target=200, min_size=60, max_size=300)
    assert len(chunks) >= 2
    assert all(len(c.text) <= 300 for c in chunks)
    assert any("考勤管理" in c.section_path for c in chunks)
    # 条文内容必须完整进入某个 chunk（背诵样本依赖它）
    assert any("第一条" in c.text for c in chunks)
    # chunk_id 稳定
    again = chunker.chunk_blocks(blocks, target=200, min_size=60, max_size=300)
    assert [c.chunk_id for c in chunks] == [c.chunk_id for c in again]


def test_assemble_three_types_and_dedup() -> None:
    chunks = [
        {"chunk_id": "c1", "section_path": "第二章", "text": "第一条 员工应当遵守工作时间，迟到三十分钟以上按旷工半日处理。"},
        {"chunk_id": "c2", "section_path": "第三章", "text": "报销应当在费用发生后三十日内提交单据。"},
    ]
    qa_records = [
        {
            "chunk_id": "c1",
            "section_path": "第二章",
            "qa": [
                {"question": "迟到多久算旷工？", "answer": "根据第一条，迟到三十分钟以上按旷工半日处理。"},
                {"question": "迟到多久算旷工?", "answer": "重复问题应被去掉"},
            ],
            "uncovered": ["迟到扣多少钱？"],
        }
    ]

    # 无 QA 记录：quote / refusal 仍可组装
    result = asm.assemble(None, chunks)
    assert result["by_type"].get("quote", 0) >= 1
    assert result["by_type"].get("refusal", 0) >= 10

    # 带 QA 记录的完整组装
    from apps.data_factory.assemble import build_qa_samples, build_refusal_samples, dedup

    qa_samples = build_qa_samples(qa_records)
    assert len(qa_samples) == 2
    merged, dropped = dedup(qa_samples)
    assert dropped == 1  # 标点差异的重复问题被去掉
    refusal = build_refusal_samples(qa_records, ["迟到扣多少钱？", "公司股票代码是多少？"])
    assert len(refusal) == 2  # uncovered 与静态池合并去重


def test_quote_samples_carry_clause_text() -> None:
    chunks = [
        {"chunk_id": "c1", "section_path": "第二章", "text": "第一条 员工应当遵守工作时间，迟到三十分钟以上按旷工半日处理。第二条 请假需提前申请。"}
    ]
    quotes = asm.build_quote_samples(chunks, max_per_chunk=2)
    assert len(quotes) == 2
    first_answer = quotes[0]["messages"][2]["content"]
    assert "第一条" in first_answer and "旷工半日" in first_answer


def test_split_grouped_no_leakage() -> None:
    def sample(i: int, chunk: str | None, stype: str = "qa") -> dict:
        return {
            "id": f"{stype}-{i}",
            "type": stype,
            "chunk_id": chunk,
            "messages": [
                {"role": "system", "content": "s"},
                {"role": "user", "content": f"q{i}"},
                {"role": "assistant", "content": "a"},
            ],
        }

    samples = [sample(i, f"c{i // 3}") for i in range(30)]  # 10 组
    samples += [sample(100 + i, None, "refusal") for i in range(6)]
    train, val, test, manifest = splitter.split_grouped(samples, (0.8, 0.1, 0.1), seed=42)

    total = len(train) + len(val) + len(test)
    assert total == len(samples)
    for split in (train, val, test):
        ids = [s["id"] for s in split]
        assert len(ids) == len(set(ids))
    get_chunk = lambda split: {s["chunk_id"] for s in split if s["chunk_id"]}  # noqa: E731
    assert not (get_chunk(train) & get_chunk(val))
    assert not (get_chunk(train) & get_chunk(test))
    assert not (get_chunk(val) & get_chunk(test))
    assert manifest["train"]["samples"] == len(train)
    assert manifest["train"]["by_type"].get("qa", 0) > 0


def test_ratios_must_sum_to_one() -> None:
    with pytest.raises(ValueError):
        splitter.split_grouped([], (0.7, 0.1, 0.1))
